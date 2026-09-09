#!/usr/bin/env python3
"""ETL материалов папки «Карасайский район» → passports/data/karasai/*.json."""

from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT.parent / "Карасайский район"
OUT = ROOT / "data" / "karasai"
sys.path.insert(0, str(ROOT / "scripts"))

from digitize_itog import digitize_itog  # noqa: E402

KASKELEN_STREETS = {
    "абылай", "жангозин", "наурызбай", "райымбек", "абая", "ерлепес",
    "заводской", "алатау", "достык", "кенесары", "новостройка",
}
IRGELI_MARKERS = {
    "иргели", "irgeli", "алтын орда", "асыл арман", "апорт", "мерей",
    "исагул", "коксай", "кемер", "ташкент",
}


def clean(t) -> str:
    return re.sub(r"\s+", " ", str(t).replace("\xa0", " ")).strip()


def article_base(qual: str) -> str:
    m = re.search(r"ст\.?\s*(\d+(?:-\d+)?)", qual or "", re.I)
    return m.group(1) if m else ""


def classify_place(text: str) -> str:
    t = (text or "").lower()
    if any(m in t for m in ("иргели", "irgeli", "алтын орда", "апорт", "асыл арман")):
        return "irgeli"
    if "коксай" in t:
        return "koksay"
    if "кемер" in t:
        return "kemertogan"
    if "каскелен" in t or "қaскелен" in t:
        return "kaskelen"
    if any(s in t for s in KASKELEN_STREETS):
        return "kaskelen"
    return "other"


def load_erdr(path: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [clean(c) for c in row]
            continue
        if not any(row):
            continue
        rec = {headers[j]: clean(row[j]) if j < len(row) and row[j] is not None else "" for j in range(len(headers))}
        rows.append(rec)
    wb.close()

    by_settlement = Counter(rec.get("31.Место совершения Населенный пункт", "") for rec in rows)
    by_qual = Counter(rec.get("10.Квалификация", "") for rec in rows)
    by_place_class = Counter(classify_place(rec.get("31.Место совершения Населенный пункт", "") + " " + rec.get("Место совершения", "")) for rec in rows)

    def count_qual(prefix: str, subset=None) -> int:
        data = subset or rows
        return sum(1 for r in data if (r.get("10.Квалификация") or "").startswith(prefix))

    kask_rows = [r for r in rows if classify_place(r.get("31.Место совершения Населенный пункт", "") + " " + r.get("Место совершения", "")) == "kaskelen"]
    irg_rows = [r for r in rows if classify_place(r.get("31.Место совершения Населенный пункт", "") + " " + r.get("Место совершения", "")) == "irgeli"]

    hotspots = {
        "altyn_orda": sum(1 for r in irg_rows if "алтын" in (r.get("Место совершения", "") + r.get("31.Место совершения Населенный пункт", "")).lower()),
        "asyl_arman": sum(1 for r in irg_rows if "асыл" in (r.get("Место совершения", "") + r.get("31.Место совершения Населенный пункт", "")).lower()),
        "aport": sum(1 for r in irg_rows if "апорт" in (r.get("Место совершения", "") + r.get("31.Место совершения Населенный пункт", "")).lower()),
        "abylay": sum(1 for r in kask_rows if "абылай" in (r.get("Место совершения", "") + r.get("31.Место совершения Населенный пункт", "")).lower()),
    }

    return {
        "total": len(rows),
        "kaskelen": len(kask_rows),
        "irgeli": len(irg_rows),
        "fraud_190": count_qual("ст.190"),
        "theft_188": count_qual("ст.188"),
        "by_settlement_top": by_settlement.most_common(12),
        "by_qual_top": by_qual.most_common(15),
        "by_class": dict(by_place_class),
        "hotspots_erdr": hotspots,
        "kaskelen_fraud": count_qual("ст.190", kask_rows),
        "irgeli_theft": count_qual("ст.188", irg_rows),
    }


def load_admin(path: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Группа отчетов об административ"]
    headers = None
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 5:
            headers = [clean(c) for c in row]
            continue
        if i < 6 or not row or not row[0]:
            continue
        rec = {headers[j]: clean(row[j]) if j < len(row) and row[j] is not None else "" for j in range(len(headers))}
        place = rec.get("2.1 Место совершения правонарушения", "") + " " + rec.get("30. Юр.адрес Насел.пункт", "")
        rec["_class"] = classify_place(place)
        rows.append(rec)
    wb.close()

    KEY = ("73", "434", "440", "442", "127", "200")

    def art_counts(subset):
        c = Counter(rec.get("9. Квалификация", "") for rec in subset)
        grouped = Counter()
        for art, n in c.items():
            base = article_base(art)
            if base:
                grouped[base] += n
        return {f"ст.{k}": grouped[k] for k in KEY if grouped[k]}

    by_class = defaultdict(list)
    for r in rows:
        by_class[r["_class"]].append(r)

    return {
        "total": len(rows),
        "by_class": {k: len(v) for k, v in by_class.items()},
        "kaskelen_key": art_counts(by_class["kaskelen"]),
        "kaskelen_street_inferred": art_counts(by_class["kaskelen"] + [r for r in by_class["other"] if classify_place(r.get("2.1 Место совершения правонарушения", "")) == "kaskelen"]),
        "irgeli_key": art_counts(by_class["irgeli"]),
        "district_key": art_counts(rows),
        "top_articles": Counter(r.get("9. Квалификация", "") for r in rows).most_common(20),
    }


def load_med_xls(path: Path, kind: str) -> dict:
    import xlrd

    wb = xlrd.open_workbook(str(path))
    sh = wb.sheet_by_index(0)
    header_row = 3 if sh.nrows > 3 and "ИИН" in clean(sh.cell_value(3, 1)) else 4
    recs = []
    for i in range(header_row + 1, sh.nrows):
        fio = clean(sh.cell_value(i, 0))
        if not fio:
            continue
        recs.append({
            "fio": fio,
            "iin": clean(sh.cell_value(i, 1)),
            "addr": clean(sh.cell_value(i, 2)),
            "mkb": clean(sh.cell_value(i, 5)),
        })
    by_loc = Counter()
    f10 = 0
    for r in recs:
        loc = classify_place(r["addr"])
        by_loc[loc] += 1
        if r["mkb"].upper().startswith("F10"):
            f10 += 1
    return {
        "kind": kind,
        "total": len(recs),
        "by_location": dict(by_loc),
        "f10_alcohol": f10,
        "kaskelen": by_loc["kaskelen"],
        "irgeli": by_loc["irgeli"],
    }


def load_commission_pdf(path: Path, session_id: str) -> dict:
    import pymupdf

    doc = pymupdf.open(str(path))
    text = "\n".join(doc[i].get_text() for i in range(len(doc)))
    doc.close()

    assignments = []
    for m in re.finditer(r"(\d+\.\d+)\.\s*(.{20,200}?)(?=\d+\.\d+\.|ШЕШ|Комиссия|$)", text, re.S):
        assignments.append({"ref": m.group(1), "text": clean(m.group(2))[:300]})

    date_m = re.search(r"(\d{1,2})\s+(?:ақпан|мамыр|феврал|мая)\s+(\d{4})", text, re.I)
    return {
        "id": session_id,
        "file": path.name,
        "pages": len(pymupdf.open(str(path))),
        "date_hint": date_m.group(0) if date_m else "",
        "assignments_count": len(assignments),
        "assignments": assignments[:25],
        "topics": {
            "lighting": bool(re.search(r"жарь|освещ|жарық", text, re.I)),
            "video": bool(re.search(r"бейне|видео|кamera", text, re.I)),
            "alcohol": bool(re.search(r"алкогол|23:00|рейд", text, re.I)),
            "minors": bool(re.search(r"кәмелет|несовершен", text, re.I)),
            "fraud": bool(re.search(r"алаяқ|мошенни", text, re.I)),
        },
        "execution_generic": bool(re.search(r"жургізілуде|ведётся|atkar", text, re.I)),
    }


def ocr_pdf_meta(path: Path) -> dict:
    import pymupdf

    doc = pymupdf.open(str(path))
    meta = {"pages": len(doc), "has_text": bool(doc[0].get_text().strip()), "images": len(doc[0].get_images())}
    if not meta["has_text"]:
        try:
            import pytesseract
            from PIL import Image
            import io as iolib

            pix = doc[0].get_pixmap(dpi=200)
            img = Image.open(iolib.BytesIO(pix.tobytes("png")))
            meta["ocr_text"] = pytesseract.image_to_string(img, lang="rus+kaz")[:2000]
        except Exception as exc:
            meta["ocr_error"] = str(exc)
    doc.close()
    return meta


def extract_passport_kpis(passport: dict) -> dict:
    admin = {a["indicator"]: a.get("count") for a in passport.get("admin_practice", [])}
    registry = {r["category"]: r.get("count") for r in passport.get("registry", [])}
    crimes = passport.get("summary", {}).get("crimes", {})
    return {
        "crimes_current": crimes.get("current"),
        "crimes_previous": crimes.get("previous"),
        "population_raw": passport.get("summary", {}).get("population_raw"),
        "rate_raw": passport.get("summary", {}).get("rate_raw"),
        "erdr_records": passport.get("summary", {}).get("erdr_records"),
        "admin_total": admin.get("Всего зарегистрировано административных правонарушений"),
        "st440": admin.get("ст.440 КоАП — распитие / появление в пьяном виде") or admin.get("ст.440 КоАП — распитие и появление в состоянии опьянения"),
        "st442": admin.get("ст.442 КоАП — несовершеннолетние ночью вне жилища"),
        "st73": admin.get("ст.73 КоАП — семейно-бытовая сфера"),
        "registry_alcohol": registry.get("Лица, злоупотребляющие алкоголем"),
        "registry_minors": registry.get("Несовершеннолетние"),
        "registry_probation": registry.get("Под пробацией"),
        "kpi_filled": sum(1 for k in passport.get("kpi_execution", []) if k.get("value")),
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    kaskelen_pp = digitize_itog(SRC / "Крим_паспорт_Каскелен_итог.docx", "kaskelen")
    irgeli_pp = digitize_itog(SRC / "Крим_паспорт_Иргели_итог.docx", "irgeli")

    (ROOT / "data" / "kaskelen.json").write_text(
        json.dumps(kaskelen_pp, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (ROOT / "data" / "irgeli.json").write_text(
        json.dumps(irgeli_pp, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    bundle = {
        "meta": {
            "district": "Карасайский район",
            "period": "7 месяцев 2026",
            "sources": [p.name for p in SRC.iterdir() if p.is_file()],
        },
        "passport_kpis": {
            "kaskelen": extract_passport_kpis(kaskelen_pp),
            "irgeli": extract_passport_kpis(irgeli_pp),
        },
        "erdr": load_erdr(SRC / "Каскелен за 2026г место совершения 3.xlsx"),
        "admin": load_admin(
            SRC / "Группа отчетов об административных правонарушениях (ф.1-АД) данные к отчету 1-АД - 28 августа 2026 г. в 16_18_44.xlsx"
        ),
        "narcology": load_med_xls(SRC / "Наркология список.xls", "narcology"),
        "psychiatry": load_med_xls(SRC / "Психиатрия список.xls", "psychiatry"),
        "commission": [
            load_commission_pdf(SRC / "Правонарушение 1 квартл.pdf", "2026-Q1-02-23"),
            load_commission_pdf(SRC / "Правонарушения 2 квартл.pdf", "2026-Q2-05-21"),
        ],
        "doc_19_356": ocr_pdf_meta(SRC / "19-356.pdf"),
        "population_alternatives": {
            "kaskelen": [87023, 84199, 89000],
            "irgeli_passport": 63152,
            "irgeli_gov_kz_2025": 43100,
        },
    }

    out_path = OUT / "bundle.json"
    out_path.write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Written {out_path}")
    print(f"Updated kaskelen.json, irgeli.json from _итог")


if __name__ == "__main__":
    main()
