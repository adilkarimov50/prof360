#!/usr/bin/env python3
"""Читаемая аналитическая справка (Word) на русском языке."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
REPORT_JSON = ROOT / "passports" / "data" / "karasai" / "report_ru.json"


def _load() -> dict:
    return json.loads(REPORT_JSON.read_text(encoding="utf-8"))


def _fmt(n) -> str:
    if n is None:
        return "—"
    if isinstance(n, int):
        return f"{n:,}".replace(",", " ")
    return str(n)


def _table(doc: Document, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)


def build_readable_docx(author: str = "") -> bytes:
    r = _load()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    run = title.add_run(r["title"])
    run.bold = True
    run.font.size = Pt(14)

    sub = doc.add_paragraph(r["subtitle"])
    sub.runs[0].italic = True

    doc.add_paragraph(
        f"Период: {r['period']}. Дата подготовки: {date.today().strftime('%d.%m.%Y')}."
        + (f" Составитель: {author}." if author else "")
    )

    doc.add_heading("Краткое резюме", level=1)
    doc.add_paragraph(r["executive_summary"])

    doc.add_heading("1. Источники и нормативная база", level=1)
    doc.add_paragraph("Источники:")
    for s in r["sources"]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph("Нормативная база:")
    for law in r["legal_basis"]:
        doc.add_paragraph(law, style="List Bullet")

    for key, num in (("kaskelen", "2"), ("irgeli", "3")):
        block = r[key]
        doc.add_heading(f"{num}. {block['name']}", level=1)
        doc.add_paragraph(block["headline"]).runs[0].bold = True
        rows = [(f["label"], f["value"], f.get("note", "")) for f in block["facts"]]
        _table(doc, ("Показатель", "Значение", "Комментарий"), rows)

    doc.add_heading("4. Матрица расхождений «на бумаге vs в данных»", level=1)
    _table(
        doc,
        ("Область", "В паспорте", "По данным", "Эффект"),
        [(d["area"], d["paper"], d["data"], d["effect"]) for d in r["discrepancies"]],
    )

    doc.add_heading("5. Профилактический учёт: воронка разрыва", level=1)
    for loc_key, loc_name in (("kaskelen", "г. Каскелен"), ("irgeli", "Иргелинский с.о.")):
        doc.add_paragraph(loc_name).runs[0].bold = True
        funnel = r["prevention_funnel"][loc_key]
        _table(doc, ("Этап", "Численность"), [(s["stage"], _fmt(s["count"])) for s in funnel])
        doc.add_paragraph("")

    doc.add_heading("6. Комиссия при акимате района (перевод с казахского)", level=1)
    for sess in r["commission"]:
        doc.add_heading(f"{sess['title']} — {sess['date']}, {sess['place']}", level=2)
        doc.add_paragraph(f"Председатель: {sess['chair']}.")
        doc.add_paragraph("Вопросы повестки:")
        for item in sess["agenda"]:
            doc.add_paragraph(item, style="List Bullet")
        doc.add_paragraph("Принятые решения:")
        for dec in sess["decisions"]:
            doc.add_paragraph(f"{dec['ref']}. {dec['text']}", style="List Number")
        doc.add_paragraph(f"Оценка исполнения: {sess['execution_note']}")

    doc.add_heading("7. Правовая оценка", level=1)
    _table(
        doc,
        ("Проблема", "Закон", "Подзаконный акт", "Акт реагирования"),
        [(x["issue"], x["law"], x["acts"], x["reaction"]) for x in r["legal_findings"]],
    )

    doc.add_heading("8. Предложения прокурору района", level=1)
    for rec in r["recommendations"]:
        doc.add_paragraph(rec, style="List Number")

    doc.add_heading("Заключение", level=1)
    doc.add_paragraph(r["conclusion"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main() -> None:
    out = ROOT / "passports" / "data" / "karasai" / "Справка_Карасай_реальная_картина.docx"
    out.write_bytes(build_readable_docx())
    print(f"Written {out}")


if __name__ == "__main__":
    main()
