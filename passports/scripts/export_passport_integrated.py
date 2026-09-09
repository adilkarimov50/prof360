#!/usr/bin/env python3
"""Встраивание интегрированного анализа (сверка профучёта) в кримпаспорта Каскелен/Иргели.

Обновляет:
  - data/kaskelen.json, data/irgeli.json — блок integrated
  - Карасайский район/Крим_паспорт_*_итог.docx — раздел 22
  - docs/assets/js/integrated_maps.js — данные для карт

Запуск после integrate_passport_crossmatch.py:
  PYTHONPATH=passports/scripts backend/.venv/bin/python3 passports/scripts/export_passport_integrated.py
"""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
KARASAI = DATA / "karasai"
OUT_DIR = ROOT.parent / "Карасайский район"
INTEGRATED_JS = ROOT / "docs" / "assets" / "js" / "integrated_maps.js"

PASSPORTS = {
    "kaskelen": {
        "json": DATA / "kaskelen.json",
        "docx": OUT_DIR / "Крим_паспорт_Каскелен_итог.docx",
        "label": "г. Каскелен",
        "markers": ("каскелен", "kaskelen", "абыл", "мошен"),
    },
    "irgeli": {
        "json": DATA / "irgeli.json",
        "docx": OUT_DIR / "Крим_паспорт_Иргели_итог.docx",
        "label": "Иргелинский с.о.",
        "markers": ("иргел", "irgeli", "алтын орда", "асыл арм", "поток"),
    },
}

SECTION_TITLE = "22. Сверка профучёта ОВД и интегрированный анализ"


def load_integrated() -> dict:
    path = KARASAI / "integrated_analysis.json"
    if not path.exists():
        raise SystemExit(f"Нет {path}. Сначала: integrate_passport_crossmatch.py")
    return json.loads(path.read_text(encoding="utf-8"))


def _matches(text: str, markers: tuple[str, ...]) -> bool:
    low = text.lower()
    return any(m in low for m in markers)


def slice_locality(locality_id: str, data: dict) -> dict:
    meta = PASSPORTS[locality_id]
    markers = meta["markers"]
    label = meta["label"]

    connections = [
        c for c in data["connections"]
        if _matches(c.get("kaskelen", "") + c.get("irgeli", "") + c.get("conclusion", ""), markers)
        or c.get("kaskelen") == "—" and locality_id == "irgeli" and _matches(c.get("irgeli", ""), markers)
        or c.get("irgeli") == "—" and locality_id == "kaskelen" and _matches(c.get("kaskelen", ""), markers)
    ]
    if len(connections) < 4:
        connections = data["connections"]

    locality_compare = [
        r for r in data["locality_compare"]
        if label.split()[0].lower() in r["locality"].lower()
        or (locality_id == "kaskelen" and "Каскелен" in r["locality"])
        or (locality_id == "irgeli" and "Иргел" in r["locality"])
    ]

    passport_vs = [
        r for r in data["passport_vs_data"]
        if _matches(" ".join(r), markers)
    ]

    problems = [
        c for c in data["conclusions"]
        if _matches(c, markers) or "район" in c.lower() or "сверка" in c.lower() or "kpi" in c.lower()
    ]
    recs = [
        r for r in data["recommendations"]
        if _matches(r, markers) or "kpi" in r.lower() or "обмен" in r.lower() or "унифиц" in r.lower()
    ]

    funnel = data["funnel"][locality_id]
    cs = data["crossmatch_summary"]

    if locality_id == "kaskelen":
        summary = (
            "Person-level сверка 3 236 лиц (августовские реестры ОВД, наркология, психиатрия, ЕРДР) "
            "подтверждает: Каскелен даёт 30,5% преступлений района при 15,4% профучёта ОВД. "
            "Воронка «переворачивается»: 87 лиц на алкопрофучёте vs 825 ст.440. "
            "Мошенничество (~54% структуры) вне реестров. 0 лиц с полным пересечением ОВД+мед+УД."
        )
        map_kpis = [
            {"value": "30,5%", "label": "преступлений района", "warn": True},
            {"value": "15,4%", "label": "профучёта ОВД", "warn": True},
            {"value": "825/87", "label": "ст.440 / учёт ОВД", "warn": True},
            {"value": "0", "label": "ОВД+мед+УД", "warn": True},
        ]
    else:
        summary = (
            "Сверка подтверждает занижение уровня преступности: при 43 100 жителей — ~51/10 тыс. (+46%). "
            "16 лиц на алкопрофучёте vs 93 ст.440. Алтын Орда (88 фактов паспорт / 26 ЕРДР) — hotspot без "
            "объектного профучёта посетителей. Дневной поток 80–90 тыс. не учтён в реестрах."
        )
        map_kpis = [
            {"value": "~51", "label": "уровень /10 тыс. (реал.)", "warn": True},
            {"value": "93/16", "label": "ст.440 / учёт ОВД", "warn": True},
            {"value": "88/26", "label": "Алтын Орда пасп/ЕРДР", "warn": True},
            {"value": "0", "label": "ОВД+мед+УД", "warn": True},
        ]

    return {
        "generated": data["meta"]["generated"],
        "title": SECTION_TITLE,
        "summary": summary,
        "connections": connections,
        "locality_compare": locality_compare,
        "passport_vs_data": passport_vs,
        "funnel": funnel,
        "crossmatch_summary": {
            "unique_persons": cs["unique_persons"],
            "match_ovd_med": cs["match_ovd_med"],
            "match_ovd_crime": cs["match_ovd_crime"],
            "match_all_three": cs["match_all_three"],
            "gap_med_no_police": cs["gap_med_no_police"],
            "gap_crime_no_police": cs["gap_crime_no_police"],
            "addr_full_linked": cs["addr_full_linked"],
        },
        "problems": problems,
        "recommendations": recs,
        "map_kpis": map_kpis,
    }


def patch_json(locality_id: str, integrated_slice: dict) -> None:
    path = PASSPORTS[locality_id]["json"]
    doc = json.loads(path.read_text(encoding="utf-8"))
    doc["integrated"] = integrated_slice
    path.write_text(json.dumps(doc, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"  JSON {path.relative_to(ROOT)}")


def _table(doc: Document, cols: int):
    tbl = doc.add_table(rows=1, cols=cols)
    try:
        tbl.style = "Table Grid"
    except KeyError:
        pass
    return tbl


def _add_para(doc: Document, text: str, *, bold: bool = False, size: float = 12) -> None:
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"


def _remove_existing_section(doc: Document) -> None:
    """Удалить ранее добавленный раздел 22, если перегенерация."""
    body = doc.element.body
    children = list(body)
    start_idx = None
    for i, child in enumerate(children):
        if not child.tag.endswith("p"):
            continue
        text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("t")).strip()
        if text.startswith("22.") and "интегрирован" in text.lower():
            start_idx = i
            break
    if start_idx is not None:
        for child in children[start_idx:]:
            body.remove(child)


def append_docx(locality_id: str, slice_data: dict) -> None:
    path = PASSPORTS[locality_id]["docx"]
    if not path.exists():
        print(f"  SKIP docx — нет файла {path}")
        return

    doc = Document(path)
    _remove_existing_section(doc)

    _add_para(doc, SECTION_TITLE, bold=True, size=13)
    _add_para(doc, f"Дата сверки: {slice_data['generated']}. Источник: августовские реестры ОВД, "
                "наркология, психиатрия, ЕРДР, особое требование/защитка (область).")
    _add_para(doc, slice_data["summary"])

    _add_para(doc, "22.1. Скрытые логические связи", bold=True)
    for c in slice_data["connections"]:
        _add_para(doc, f"{c['id']}. {c['title']}", bold=True)
        if locality_id == "kaskelen":
            detail = c["kaskelen"]
        else:
            detail = c["irgeli"]
        if detail and detail != "—":
            _add_para(doc, detail)
        _add_para(doc, f"Вывод: {c['conclusion']}")

    _add_para(doc, "22.2. Доли населённого пункта от районного числа", bold=True)
    tbl = _table(doc, 5)
    for i, h in enumerate(("Источник", "НП", "Число", "Всего по району", "Доля")):
        tbl.rows[0].cells[i].text = h
    for row in slice_data["locality_compare"]:
        cells = tbl.add_row().cells
        cells[0].text = row["source"]
        cells[1].text = row["locality"]
        cells[2].text = str(row["count"])
        cells[3].text = str(row["district_total"])
        cells[4].text = row["share"]

    _add_para(doc, "22.3. Воронка профилактики vs административная практика", bold=True)
    t2 = _table(doc, 2)
    t2.rows[0].cells[0].text = "Этап"
    t2.rows[0].cells[1].text = "Число"
    for step in slice_data["funnel"]:
        r = t2.add_row().cells
        r[0].text = step["stage"]
        r[1].text = str(step["count"])

    cs = slice_data["crossmatch_summary"]
    _add_para(doc, "22.4. Сверка профучёта (район, person-level)", bold=True)
    t3 = _table(doc, 2)
    t3.rows[0].cells[0].text = "Показатель"
    t3.rows[0].cells[1].text = "Значение"
    for label, key in (
        ("Уникальных лиц в сверке", "unique_persons"),
        ("ОВД + мед", "match_ovd_med"),
        ("ОВД + УД", "match_ovd_crime"),
        ("ОВД + мед + УД", "match_all_three"),
        ("Медучёт без ОВД", "gap_med_no_police"),
        ("Подозреваемые без ОВД", "gap_crime_no_police"),
        ("Адресных связей (проживание)", "addr_full_linked"),
    ):
        r = t3.add_row().cells
        r[0].text = label
        r[1].text = f"{cs[key]:,}".replace(",", " ")

    if slice_data["passport_vs_data"]:
        _add_para(doc, "22.5. Паспорт vs оперативные данные", bold=True)
        t4 = _table(doc, 4)
        for i, h in enumerate(("НП", "Показатель", "Значения", "Эффект")):
            t4.rows[0].cells[i].text = h
        for row in slice_data["passport_vs_data"]:
            cells = t4.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)

    _add_para(doc, "22.6. Выявленные проблемы (новые выводы)", bold=True)
    for i, prob in enumerate(slice_data["problems"], 1):
        _add_para(doc, f"{i}. {prob}")

    _add_para(doc, "22.7. Рекомендации", bold=True)
    for i, rec in enumerate(slice_data["recommendations"], 1):
        _add_para(doc, f"{i}. {rec}")

    doc.save(path)
    print(f"  DOCX {path.name}")


def write_integrated_js(slices: dict) -> None:
    payload = {k: v for k, v in slices.items()}
    INTEGRATED_JS.parent.mkdir(parents=True, exist_ok=True)
    body = json.dumps(payload, ensure_ascii=False, indent=1)
    INTEGRATED_JS.write_text(
        "/* Сформировано export_passport_integrated.py — не редактировать вручную. */\n"
        f"window.INTEGRATED_MAPS = {body};\n",
        encoding="utf-8",
    )
    print(f"  JS {INTEGRATED_JS.relative_to(ROOT)}")


def main() -> None:
    data = load_integrated()
    slices = {}
    for loc_id in PASSPORTS:
        print(f"→ {PASSPORTS[loc_id]['label']}")
        sl = slice_locality(loc_id, data)
        slices[loc_id] = sl
        patch_json(loc_id, sl)
        append_docx(loc_id, sl)
    write_integrated_js(slices)
    print("Готово.")


if __name__ == "__main__":
    main()
