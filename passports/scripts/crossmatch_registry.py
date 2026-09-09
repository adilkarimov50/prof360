#!/usr/bin/env python3
"""Полная сверка профучёта ОВД со всеми списками + адреса + уголовные дела."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT.parent / "Карасайский район"
OUT = ROOT / "data" / "karasai" / "registry_crossmatch.json"
HTML = ROOT / "docs" / "registry_crossmatch.html"

POLICE_AUGUST = {
    "ОП (общий)": "ОП август.xlsx",
    "Наркологический учёт ОВД": "нл общий август.xlsx",
    "Ранее судимые": "ранее судимый август.xlsx",
    "Адм. надзор": "адм надзор август.xlsx",
    "Особое требование": "особое требование август.xlsx",
    "Защитное предписание": "защитное предписание август.xlsx",
    "УДО": "УДОавгуст.xlsx",
}

POLICE_FULL = {
    "Особое требование (область)": "особое треб.xlsx",
    "Защитное предписание (область)": "защитка.xlsx",
}

# Семантика адресов по источникам
ADDR_RESIDENCE = "адрес фактического проживания"
ADDR_MED = "адрес по мед. учёту (мекенжай)"
ADDR_CRIME = "место совершения (нас. пункт)"
ADDR_CRIME_DIST = "район совершения преступления"
ADDR_FABULA = "адрес/место в фабуле (текст, не структурировано)"
ADDR_ORDER = "адрес фактического проживания (основание ограничения)"

LOCALITIES = {
    "kaskelen": {"label": "г. Каскелен", "tokens": ("КАСКЕЛЕН",)},
    "irgeli": {"label": "с.о. Иргели (вкл. Казмис)", "tokens": ("ИРГЕЛ", "КАЗМИС")},
}


def clean(t) -> str:
    if t is None:
        return ""
    return re.sub(r"\s+", " ", str(t).replace("\xa0", " ")).strip()


def norm_iin(v) -> str:
    s = re.sub(r"\D", "", str(v or ""))
    if len(s) != 12:
        return ""
    # Формат ИИН: ГГММДД + 6 цифр; месяц — позиции 2–3
    if not (s[2:4].isdigit() and 1 <= int(s[2:4]) <= 12):
        return ""
    if not (s[4:6].isdigit() and 1 <= int(s[4:6]) <= 31):
        return ""
    return s


def norm_text(s: str) -> str:
    s = clean(s).upper()
    for a, b in (("Ё", "Е"), ("І", "И"), ("Ұ", "У"), ("Ү", "У"), ("Қ", "К"), ("Ғ", "Г"), ("Ә", "А"), ("Ө", "О"), ("Ң", "Н")):
        s = s.replace(a, b)
    s = re.sub(r"[^A-ZА-Я0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_fio(surname, name, patronymic=None) -> str:
    return norm_text(" ".join(x for x in (surname, name, patronymic) if clean(x)))


def fio_sim(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def classify_locality(text: str) -> str | None:
    u = norm_text(text)
    if not u:
        return None
    for key, meta in LOCALITIES.items():
        if any(t in u for t in meta["tokens"]):
            return key
    return None


def is_karasai_district(text: str) -> bool:
    return "КАРАСАЙ" in norm_text(text)


def pct_str(part: int, whole: int) -> str:
    if not whole:
        return "0%"
    return f"{100 * part / whole:.1f}%"


def locality_stats(records: list[dict], text_fn, district_fn=None) -> dict:
    """Считает Каскелен/Иргели и долю от районного total."""
    if district_fn:
        pool = [r for r in records if district_fn(r)]
    else:
        pool = records
    total = len(pool)
    out: dict[str, dict] = {}
    for key, meta in LOCALITIES.items():
        cnt = sum(1 for r in pool if classify_locality(text_fn(r)) == key)
        out[key] = {
            "label": meta["label"],
            "count": cnt,
            "district_total": total,
            "share": pct_str(cnt, total),
        }
    out["_district_total"] = total
    return out


def build_address(settlement, street, house, block="", apt="", district="", region="") -> str:
    parts = [clean(region), clean(district), clean(settlement), clean(street), clean(house), clean(block), clean(apt)]
    return norm_text(" ".join(p for p in parts if p))


def addr_key(addr: str) -> str:
    """Ключ для нечёткого сравнения адресов."""
    a = norm_text(addr)
    for w in ("УЛИЦА", "КВАРТАЛ", "ПОСЕЛОК", "СЕЛО", "АУЛ", "ОБЛАСТЬ", "РАЙОН", "ДОМ", "КВ", "БЛОК", "ДРУГОЕ"):
        a = a.replace(w, " ")
    tokens = sorted(set(t for t in a.split() if len(t) > 2))
    return " ".join(tokens)


def addr_sim(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    ka, kb = addr_key(a), addr_key(b)
    if ka == kb:
        return 1.0
    return SequenceMatcher(None, ka, kb).ratio()


def load_police_august(path: Path, category: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["data"]
    out: list[dict] = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not any(row):
            continue
        iin = norm_iin(row[4])
        fio = norm_fio(row[5], row[6], row[7])
        if not iin and not fio:
            continue
        addr = build_address(
            row[16] if len(row) > 16 else "",
            row[18] if len(row) > 18 else "",
            row[19] if len(row) > 19 else "",
            row[20] if len(row) > 20 else "",
            row[21] if len(row) > 21 else "",
            row[14] if len(row) > 14 else "",
            row[11] if len(row) > 11 else "",
        )
        out.append({
            "iin": iin,
            "fio": fio,
            "dob": clean(row[8]) if len(row) > 8 else "",
            "address": addr,
            "address_type": ADDR_RESIDENCE,
            "addr_key": addr_key(addr),
            "settlement": clean(row[16]) if len(row) > 16 else "",
            "district": clean(row[14]) if len(row) > 14 else "",
            "category": clean(row[22])[:80] if len(row) > 22 else category,
            "registry_type": category,
            "status": clean(row[1]),
        })
    wb.close()
    return out


def load_police_full(path: Path, category: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["data"]
    out: list[dict] = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        addr = build_address(
            row[28] if len(row) > 28 else "",
            row[30] if len(row) > 30 else "",
            row[31] if len(row) > 31 else "",
            row[32] if len(row) > 32 else "",
            row[33] if len(row) > 33 else "",
            row[25] if len(row) > 25 else "",
            row[23] if len(row) > 23 else "",
        )
        if not addr:
            continue
        out.append({
            "iin": "",
            "fio": "",
            "address": addr,
            "address_type": ADDR_ORDER,
            "addr_key": addr_key(addr),
            "settlement": clean(row[28]) if len(row) > 28 else "",
            "district": clean(row[24]) if len(row) > 24 else "",
            "category": clean(row[2])[:80] if len(row) > 2 else category,
            "registry_type": category,
            "status": clean(row[1]),
            "doc_no": clean(row[8]) if len(row) > 8 else "",
            "rup": clean(row[16]) if len(row) > 16 else "",
        })
    wb.close()
    return out


def load_med_xls(path: Path, kind: str) -> list[dict]:
    import xlrd

    wb = xlrd.open_workbook(str(path))
    sh = wb.sheet_by_index(0)
    out: list[dict] = []
    for i in range(4, sh.nrows):
        fio_raw = clean(sh.cell_value(i, 0))
        iin = norm_iin(sh.cell_value(i, 1))
        if not iin and not fio_raw:
            continue
        addr = clean(sh.cell_value(i, 2))
        out.append({
            "iin": iin,
            "fio": norm_text(fio_raw),
            "fio_raw": fio_raw,
            "address": norm_text(addr),
            "address_type": ADDR_MED,
            "addr_key": addr_key(addr),
            "dob": clean(sh.cell_value(i, 3)),
            "mkb": clean(sh.cell_value(i, 5)),
            "source": kind,
        })
    return out


def load_suspects(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        iin = norm_iin(row[16] if len(row) > 16 else "")
        fio = norm_fio(row[13], row[14], row[15]) if len(row) > 15 else ""
        if not iin and not fio:
            continue
        crime_region = clean(row[5]) if len(row) > 5 else ""
        crime_district = clean(row[6]) if len(row) > 6 else ""
        crime_settlement = clean(row[25]) if len(row) > 25 else ""
        fabula = clean(row[7]) if len(row) > 7 else ""
        out.append({
            "iin": iin,
            "fio": fio,
            "dob": clean(row[17]) if len(row) > 17 else "",
            "erdr": clean(row[1]),
            "qual": clean(row[8]),
            "crime_region": crime_region,
            "crime_district": crime_district,
            "crime_settlement": norm_text(crime_settlement),
            "crime_place": norm_text(f"{crime_region} {crime_district} {crime_settlement}"),
            "fabula": fabula[:300],
            "occupation": clean(row[20]) if len(row) > 20 else "",
            # legacy alias — это НЕ адрес проживания
            "place": crime_settlement,
        })
    wb.close()
    return out


def index_by_iin(records: list[dict]) -> dict[str, list[dict]]:
    idx: dict[str, list[dict]] = defaultdict(list)
    for r in records:
        if r.get("iin"):
            idx[r["iin"]].append(r)
    return idx


def best_fio_match(target: str, candidates: list[dict], threshold: float = 0.88) -> dict | None:
    best, score = None, 0.0
    for c in candidates:
        s = fio_sim(target, c.get("fio", ""))
        if s > score:
            score, best = s, c
    if best and score >= threshold:
        return {**best, "fio_score": round(score, 3)}
    return None


def match_address(record: dict, addr_index: list[dict], threshold: float = 0.82) -> list[dict]:
    hits = []
    key = record.get("addr_key", "")
    for ref in addr_index:
        if key and ref.get("addr_key") == key:
            hits.append({**ref, "addr_score": 1.0})
            continue
        sc = addr_sim(record.get("address", ""), ref.get("address", ""))
        if sc >= threshold:
            hits.append({**ref, "addr_score": round(sc, 3)})
    hits.sort(key=lambda x: -x["addr_score"])
    return hits[:5]


def build_crossmatch() -> dict:
    police: list[dict] = []
    for cat, fn in POLICE_AUGUST.items():
        p = SRC / fn
        if p.exists():
            police.extend(load_police_august(p, cat))

    full_registries: list[dict] = []
    for cat, fn in POLICE_FULL.items():
        p = SRC / fn
        if p.exists():
            full_registries.extend(load_police_full(p, cat))

    narco = load_med_xls(SRC / "Наркология список.xls", "наркология")
    psych = load_med_xls(SRC / "Психиатрия список.xls", "психиатрия")

    suspects_path = SRC / "Портрет подозрика последний.xlsx"
    suspects = load_suspects(suspects_path) if suspects_path.exists() else []

    narco_iin = index_by_iin(narco)
    psych_iin = index_by_iin(psych)
    suspect_iin = index_by_iin(suspects)

    # Уникальные лица по ИИН (+ без ИИН по ФИО)
    persons: dict[str, dict] = {}

    def ensure_person(rec: dict, source: str) -> str:
        key = f"iin:{rec['iin']}" if rec.get("iin") else f"fio:{rec.get('fio','')}"
        if key not in persons:
            persons[key] = {
                "iin": rec.get("iin", ""),
                "fio": rec.get("fio", ""),
                "dob": rec.get("dob", ""),
                "address_residence": "",
                "address_med": "",
                "address_crime": "",
                "fabula": "",
                "settlement": "",
                "locality": "",
                "police_types": [],
                "narco": False,
                "psych": False,
                "suspect": False,
                "erdr_list": [],
                "flags": [],
                "status": "",
            }
        return key

    police_addr_index = [r for r in police if r.get("address")]

    for r in police:
        key = ensure_person(r, "police")
        p = persons[key]
        if r["registry_type"] not in p["police_types"]:
            p["police_types"].append(r["registry_type"])
        if not p["address_residence"] and r.get("address"):
            p["address_residence"] = r["address"]
            p["settlement"] = r.get("settlement", "")
            p["locality"] = classify_locality(f"{r.get('settlement', '')} {r.get('address', '')}") or ""
        if not p["fio"]:
            p["fio"] = r["fio"]
        if not p["dob"]:
            p["dob"] = r.get("dob", "")

    for r in suspects:
        key = ensure_person(r, "suspect")
        p = persons[key]
        p["suspect"] = True
        erdr_item = {
            "erdr": r["erdr"],
            "qual": r["qual"],
            "crime_settlement": r.get("crime_settlement", ""),
            "crime_district": r.get("crime_district", ""),
            "crime_place": r.get("crime_place", ""),
        }
        if r["erdr"] and r["erdr"] not in [e["erdr"] for e in p["erdr_list"]]:
            p["erdr_list"].append(erdr_item)
        if r.get("crime_place") and not p["address_crime"]:
            p["address_crime"] = r["crime_place"]
        if r.get("fabula") and not p["fabula"]:
            p["fabula"] = r["fabula"][:200]
        if not p["fio"]:
            p["fio"] = r["fio"]

    # Медучёт: добавить лиц без полицейского учёта
    for r in narco:
        if not r["iin"]:
            continue
        key = f"iin:{r['iin']}"
        if key not in persons:
            persons[key] = {
                "iin": r["iin"],
                "fio": r["fio"],
                "dob": r.get("dob", ""),
                "address_residence": "",
                "address_med": r.get("address", ""),
                "address_crime": "",
                "fabula": "",
                "settlement": "",
                "locality": classify_locality(r.get("address", "")) or "",
                "police_types": [],
                "narco": True,
                "psych": False,
                "suspect": False,
                "erdr_list": [],
                "flags": ["только наркология"],
                "status": "РАСХОЖДЕНИЕ: мед. без ОВД",
            }
        else:
            persons[key]["narco"] = True
            if r.get("address"):
                persons[key]["address_med"] = r["address"]

    for r in psych:
        if not r["iin"]:
            continue
        key = f"iin:{r['iin']}"
        if key not in persons:
            persons[key] = {
                "iin": r["iin"],
                "fio": r["fio"],
                "dob": r.get("dob", ""),
                "address_residence": "",
                "address_med": r.get("address", ""),
                "address_crime": "",
                "fabula": "",
                "settlement": "",
                "locality": classify_locality(r.get("address", "")) or "",
                "police_types": [],
                "narco": False,
                "psych": True,
                "suspect": False,
                "erdr_list": [],
                "flags": ["только психиатрия"],
                "status": "РАСХОЖДЕНИЕ: мед. без ОВД",
            }
        else:
            persons[key]["psych"] = True

    # Обогащение флагами
    table_rows: list[dict] = []
    for key, p in persons.items():
        iin = p["iin"]
        flags = list(p.get("flags", []))

        if iin and iin in narco_iin:
            p["narco"] = True
            p["narco_mkb"] = narco_iin[iin][0].get("mkb", "")
            if "наркология" not in flags:
                flags.append("наркология")
        elif p["fio"]:
            m = best_fio_match(p["fio"], narco)
            if m:
                p["narco"] = True
                p["narco_mkb"] = m.get("mkb", "")
                flags.append("наркология (ФИО)")

        if iin and iin in psych_iin:
            p["psych"] = True
            p["psych_mkb"] = psych_iin[iin][0].get("mkb", "")
            if "психиатрия" not in flags:
                flags.append("психиатрия")
        elif p["fio"]:
            m = best_fio_match(p["fio"], psych)
            if m:
                p["psych"] = True
                p["psych_mkb"] = m.get("mkb", "")
                flags.append("психиатрия (ФИО)")

        if iin and iin in suspect_iin:
            p["suspect"] = True
            for s in suspect_iin[iin]:
                item = {
                    "erdr": s["erdr"],
                    "qual": s["qual"],
                    "crime_settlement": s.get("crime_settlement", ""),
                    "crime_district": s.get("crime_district", ""),
                    "crime_place": s.get("crime_place", ""),
                }
                if s["erdr"] not in [e["erdr"] for e in p["erdr_list"]]:
                    p["erdr_list"].append(item)
                if s.get("crime_place") and not p["address_crime"]:
                    p["address_crime"] = s["crime_place"]
                if s.get("fabula") and not p["fabula"]:
                    p["fabula"] = s["fabula"][:200]
            flags.append("уголовное дело")

        # Сверка только адресов проживания: ОВД ↔ мед
        addr_med = p.get("address_med") or (narco_iin.get(iin, [{}])[0].get("address") if iin else "")
        if not addr_med and iin and iin in psych_iin:
            addr_med = psych_iin[iin][0].get("address", "")
        addr_match = ""
        if p.get("address_residence") and addr_med:
            sc = addr_sim(p["address_residence"], addr_med)
            if sc >= 0.75:
                addr_match = f"проживание совпадает ({sc:.0%})"
                flags.append("адрес прожив.: ОВД≈мед")
            elif sc >= 0.5:
                addr_match = f"проживание частично ({sc:.0%})"
                flags.append("адрес прожив.: частично")
            else:
                addr_match = f"проживание разное ({sc:.0%})"
                flags.append("адрес прожив.: расхождение")

        # Статус
        has_police = bool(p["police_types"])
        has_med = p["narco"] or p["psych"]
        has_crime = p["suspect"] or bool(p["erdr_list"])

        if has_police and has_med and has_crime:
            status = "СОВПАДЕНИЕ: ОВД + мед + УД"
        elif has_police and has_med:
            status = "СОВПАДЕНИЕ: ОВД + мед"
        elif has_police and has_crime:
            status = "СОВПАДЕНИЕ: ОВД + уголовное дело"
        elif has_med and has_crime and not has_police:
            status = "РАСХОЖДЕНИЕ: мед + УД, нет ОВД"
        elif has_med and not has_police:
            status = "РАСХОЖДЕНИЕ: мед. без ОВД"
        elif has_police and not has_med:
            status = "РАСХОЖДЕНИЕ: ОВД без мед."
        elif has_crime and not has_police:
            status = "РАСХОЖДЕНИЕ: УД без профучёта"
        else:
            status = "нет связей"

        p["flags"] = flags
        p["status"] = status
        p["addr_match"] = addr_match

        table_rows.append({
            "fio": p["fio"],
            "iin": p["iin"],
            "dob": p.get("dob", ""),
            "settlement": p.get("settlement", ""),
            "locality": LOCALITIES.get(p.get("locality", ""), {}).get("label", p.get("locality", "") or "—"),
            "address_residence": p.get("address_residence", "")[:100],
            "address_med": (addr_med or "")[:100],
            "address_crime": (p.get("address_crime") or "")[:100],
            "addr_match": addr_match,
            "police_types": ", ".join(p["police_types"]) if p["police_types"] else "—",
            "narco": "да" if p["narco"] else "нет",
            "narco_mkb": p.get("narco_mkb", ""),
            "psych": "да" if p["psych"] else "нет",
            "psych_mkb": p.get("psych_mkb", ""),
            "suspect": "да" if p["suspect"] else "нет",
            "erdr": "; ".join(f"{e['erdr']} ({e['qual']})" for e in p["erdr_list"][:2]) or "—",
            "crime_place": (p["erdr_list"][0].get("crime_settlement", "") if p["erdr_list"] else "") or "—",
            "flags": ", ".join(flags) if flags else "—",
            "status": status,
        })

    # Адресная привязка: только «проживание» (областной реестр ↔ ОВД)
    addr_links: list[dict] = []
    for fr in full_registries:
        if not is_karasai_district(f"{fr.get('district', '')} {fr.get('address', '')}"):
            continue
        hits = match_address(fr, police_addr_index)
        for h in hits[:1]:
            addr_links.append({
                "full_registry": fr["registry_type"],
                "address": fr["address"][:100],
                "address_type": ADDR_ORDER,
                "settlement": fr.get("settlement", ""),
                "matched_fio": h.get("fio", ""),
                "matched_iin": h.get("iin", ""),
                "matched_police_type": h.get("registry_type", ""),
                "matched_address_type": ADDR_RESIDENCE,
                "addr_score": h.get("addr_score", 0),
                "doc_no": fr.get("doc_no", ""),
                "locality": classify_locality(f"{fr.get('settlement', '')} {fr.get('address', '')}") or "",
            })

    # Suspect без ОВД
    suspect_no_police = [
        {
            "fio": s["fio"],
            "iin": s["iin"],
            "erdr": s["erdr"],
            "qual": s["qual"],
            "crime_settlement": s.get("crime_settlement", ""),
            "crime_district": s.get("crime_district", ""),
            "crime_place": s.get("crime_place", ""),
            "address_type": ADDR_CRIME,
        }
        for s in suspects
        if s["iin"] and s["iin"] not in {p["iin"] for p in persons.values() if p.get("police_types")}
    ]

    karasai_full_reg = [r for r in full_registries if is_karasai_district(f"{r.get('district', '')} {r.get('address', '')}")]

    localities = {
        "ovd_residence": locality_stats(police, lambda r: f"{r.get('settlement', '')} {r.get('address', '')}"),
        "med_narco_residence": locality_stats(narco, lambda r: r.get("address", "")),
        "med_psych_residence": locality_stats(psych, lambda r: r.get("address", "")),
        "crime_location": locality_stats(
            suspects,
            lambda r: r.get("crime_settlement", "") or r.get("crime_place", ""),
            district_fn=lambda r: is_karasai_district(r.get("crime_district", "")),
        ),
        "order_residence": locality_stats(
            karasai_full_reg,
            lambda r: f"{r.get('settlement', '')} {r.get('address', '')}",
        ),
    }

    locality_table = []
    for source_key, label in (
        ("ovd_residence", "Профучёт ОВД — проживание"),
        ("med_narco_residence", "Наркология — проживание"),
        ("med_psych_residence", "Психиатрия — проживание"),
        ("crime_location", "ЕРДР — место совершения"),
        ("order_residence", "Особое треб./защитка — проживание (район)"),
    ):
        block = localities[source_key]
        total = block["_district_total"]
        for loc in ("kaskelen", "irgeli"):
            item = block[loc]
            locality_table.append({
                "source": label,
                "locality": item["label"],
                "count": item["count"],
                "district_total": total,
                "share": item["share"],
            })

    table_rows.sort(key=lambda x: (
        0 if "СОВПАДЕНИЕ" in x["status"] else 1,
        -len(x["flags"].split(",")),
        x["fio"],
    ))

    summary = {
        "police_august_records": len(police),
        "police_unique_iin": len({r["iin"] for r in police if r["iin"]}),
        "police_types": dict(Counter(r["registry_type"] for r in police)),
        "full_registry_records": len(full_registries),
        "full_registry_types": dict(Counter(r["registry_type"] for r in full_registries)),
        "narco_total": len(narco),
        "psych_total": len(psych),
        "suspects_total": len(suspects),
        "suspects_with_iin": len({s["iin"] for s in suspects if s["iin"]}),
        "unique_persons": len(persons),
        "match_ovd_med": sum(1 for r in table_rows if "ОВД + мед" in r["status"]),
        "match_ovd_crime": sum(1 for r in table_rows if "ОВД + уголовное" in r["status"] or "ОВД + мед + УД" in r["status"]),
        "match_all_three": sum(1 for r in table_rows if "ОВД + мед + УД" in r["status"]),
        "gap_med_no_police": sum(1 for r in table_rows if "мед. без ОВД" in r["status"] or "мед + УД, нет ОВД" in r["status"]),
        "gap_crime_no_police": len(suspect_no_police),
        "addr_ovd_med_match": sum(1 for r in table_rows if "адрес прожив.: ОВД≈мед" in r["flags"]),
        "addr_full_linked": len(addr_links),
        "suspect_in_police": sum(1 for s in suspects if s["iin"] in {r["iin"] for r in police if r["iin"]}),
        "karasai_full_registry": len(karasai_full_reg),
    }

    address_legend = {
        ADDR_RESIDENCE: "Списки ОВД (август): колонка «Адрес фактического проживания»",
        ADDR_MED: "Наркология/психиатрия: «Мекенжай / Адрес» — адрес учёта пациента",
        ADDR_CRIME: "Портрет подозреваемого: «31. Место совершения — населённый пункт» (НЕ проживание)",
        ADDR_CRIME_DIST: "Портрет подозреваемого: «3. Район совершения»",
        ADDR_FABULA: "Портрет подозреваемого: «9.1 Описание» — текст фабулы, адреса не структурированы",
        ADDR_ORDER: "особое треб.xlsx / защитка.xlsx: «Адрес фактического проживания» лица по ограничению",
    }

    return {
        "meta": {
            "generated": "2026-09-07",
            "district": "Карасайский район + областные реестры",
            "sources": list(POLICE_AUGUST.keys()) + list(POLICE_FULL.keys()) + ["наркология", "психиатрия", "портрет подозреваемого"],
        },
        "address_legend": address_legend,
        "localities": localities,
        "locality_table": locality_table,
        "summary": summary,
        "table": table_rows,
        "addr_links": sorted(addr_links, key=lambda x: -x["addr_score"])[:200],
        "suspect_no_police": suspect_no_police[:100],
        "multi_registry": _multi_registry(police),
    }


def _multi_registry(police: list[dict]) -> list[dict]:
    by_iin: dict[str, set] = defaultdict(set)
    fio_map: dict[str, str] = {}
    for r in police:
        if r["iin"]:
            by_iin[r["iin"]].add(r["registry_type"])
            fio_map[r["iin"]] = r["fio"]
    out = []
    for iin, types in by_iin.items():
        if len(types) > 1:
            out.append({"iin": iin, "fio": fio_map[iin], "types": sorted(types)})
    return sorted(out, key=lambda x: -len(x["types"]))[:40]


def write_html(data: dict, path: Path) -> None:
    s = data["summary"]
    esc = lambda t: str(t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def rows(items, cols):
        return "".join(
            "<tr>" + "".join(f"<td>{esc(r.get(c,''))}</td>" for c in cols) + "</tr>"
            for r in items
        )

    tbl_cols = ["fio", "iin", "locality", "police_types", "narco", "psych", "suspect", "crime_place", "addr_match", "address_residence", "address_crime", "status"]
    tbl_hdr = ["ФИО", "ИИН", "НП", "Учёт ОВД", "Нарк.", "Псих.", "УД", "Место соверш.", "Сверка прожив.", "Проживание", "Место преступ.", "Статус"]

    addr_cols = ["full_registry", "address_type", "settlement", "address", "matched_fio", "matched_police_type", "addr_score"]
    addr_hdr = ["Реестр", "Тип адреса", "НП", "Адрес проживания", "ФИО (ОВД)", "Учёт ОВД", "Схожесть"]

    loc_rows = data.get("locality_table", [])
    loc_html = "".join(
        f"<tr><td>{esc(r['source'])}</td><td>{esc(r['locality'])}</td>"
        f"<td>{r['count']}</td><td>{r['district_total']}</td><td><strong>{esc(r['share'])}</strong></td></tr>"
        for r in loc_rows
    )

    legend = data.get("address_legend", {})
    legend_html = "".join(f"<li><strong>{esc(k)}</strong> — {esc(v)}</li>" for k, v in legend.items())

    html = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Полная сверка профучёта · Карасайский район</title>
<link href="https://fonts.googleapis.com/css2?family=Source+Sans+3:wght@400;600;700&display=swap" rel="stylesheet">
<style>
body{{font-family:'Source Sans 3',sans-serif;margin:0;background:#f5f7fa;color:#1a1f2e;line-height:1.5}}
.wrap{{max-width:1280px;margin:0 auto;padding:28px 16px 60px}}
h1{{font-size:1.5rem;margin:0 0 6px}} .sub{{color:#5c6578;margin-bottom:20px}}
.kpi{{display:grid;grid-template-columns:repeat(auto-fit,minmax(130px,1fr));gap:8px;margin:18px 0 28px}}
.kpi div{{background:#fff;border:1px solid #dde2ea;border-radius:8px;padding:12px}}
.kpi .v{{font-size:1.35rem;font-weight:700}} .kpi .l{{font-size:10px;color:#5c6578;text-transform:uppercase}}
.tabs{{display:flex;gap:4px;flex-wrap:wrap;margin-bottom:16px}}
.tab{{padding:8px 14px;border:1px solid #dde2ea;background:#fff;cursor:pointer;font:inherit;border-radius:6px}}
.tab.on{{background:#1b4d8c;color:#fff;border-color:#1b4d8c}}
.panel{{display:none}} .panel.on{{display:block}}
table{{width:100%;border-collapse:collapse;font-size:13px;background:#fff;border:1px solid #dde2ea;margin-bottom:24px}}
th,td{{padding:8px 10px;border-bottom:1px solid #eef1f5;text-align:left;vertical-align:top}}
th{{background:#eef3fa;font-size:11px;text-transform:uppercase;color:#5c6578;position:sticky;top:0}}
tr.match{{background:#f0faf4}} tr.gap{{background:#fff8f6}}
.mono{{font-family:monospace;font-size:11px}}
.note{{background:#fff8e6;border-left:4px solid #c08a1e;padding:12px 16px;margin:16px 0;font-size:14px}}
.scroll{{overflow-x:auto;max-height:70vh;overflow-y:auto}}
.filter{{margin-bottom:12px}} .filter input{{padding:8px 12px;width:280px;border:1px solid #dde2ea;border-radius:6px;font:inherit}}
</style></head><body><div class="wrap">
<h1>Полная сверка подучётных списков</h1>
<p class="sub">ОВД (август) + особое требование + защитка (область) + наркология + психиатрия + портрет подозреваемого · {s['unique_persons']} лиц</p>
<div class="kpi">
<div><div class="v">{s['police_august_records']}</div><div class="l">Записей ОВД</div></div>
<div><div class="v">{s['full_registry_records']}</div><div class="l">Областные реестры</div></div>
<div><div class="v">{s['suspects_with_iin']}</div><div class="l">Подозреваемых</div></div>
<div><div class="v">{s['match_ovd_med']}</div><div class="l">ОВД + мед</div></div>
<div><div class="v">{s['match_all_three']}</div><div class="l">ОВД+мед+УД</div></div>
<div><div class="v">{s['suspect_in_police']}</div><div class="l">УД в профучёте</div></div>
<div><div class="v">{s['gap_med_no_police']}</div><div class="l">Мед без ОВД</div></div>
<div><div class="v">{s['gap_crime_no_police']}</div><div class="l">УД без ОВД</div></div>
<div><div class="v">{s['addr_full_linked']}</div><div class="l">Адрес. связей</div></div>
</div>
<div class="note">Новые списки: <strong>особое треб.xlsx</strong> ({s['full_registry_types'].get('Особое требование (область)',0)} зап.), 
<strong>защитка.xlsx</strong> ({s['full_registry_types'].get('Защитное предписание (область)',0)} зап.), 
<strong>Портрет подозреваемого</strong> ({s['suspects_total']} дел, {s['suspects_with_iin']} с ИИН).</div>
<div class="note">Сверка адресов: сопоставляются только <strong>адреса проживания</strong> (ОВД ↔ мед ↔ особое треб./защитка). 
Место совершения преступления из ЕРДР — отдельная колонка, с проживанием не смешивается.</div>
<div class="note"><ul style="margin:0;padding-left:18px">{legend_html}</ul></div>
<div class="tabs">
<button class="tab on" data-tab="all">Вся таблица</button>
<button class="tab" data-tab="match">Совпадения</button>
<button class="tab" data-tab="gap">Расхождения</button>
<button class="tab" data-tab="locality">Каскелен / Иргели</button>
<button class="tab" data-tab="addr">Адреса прожив.</button>
<button class="tab" data-tab="crime">УД без ОВД</button>
</div>
<div class="panel on" id="panel-all">
<div class="filter"><input id="search" placeholder="Поиск по ФИО, ИИН, адресу…"></div>
<div class="scroll"><table id="main-table"><thead><tr>{"".join(f"<th>{h}</th>" for h in tbl_hdr)}</tr></thead>
<tbody>{"".join(f'<tr class="{"match" if "СОВПАДЕНИЕ" in r["status"] else "gap" if "РАСХОЖДЕНИЕ" in r["status"] else ""}">' + "".join(f'<td class="{"mono" if c=="iin" else ""}">{esc(r.get(c,""))}</td>' for c in tbl_cols) + "</tr>" for r in data["table"])}</tbody></table></div>
</div>
<div class="panel" id="panel-match"><div class="scroll"><table><thead><tr>{"".join(f"<th>{h}</th>" for h in tbl_hdr)}</tr></thead>
<tbody>{rows([r for r in data["table"] if "СОВПАДЕНИЕ" in r["status"]], tbl_cols)}</tbody></table></div></div>
<div class="panel" id="panel-gap"><div class="scroll"><table><thead><tr>{"".join(f"<th>{h}</th>" for h in tbl_hdr)}</tr></thead>
<tbody>{rows([r for r in data["table"] if "РАСХОЖДЕНИЕ" in r["status"]], tbl_cols)}</tbody></table></div></div>
<div class="panel" id="panel-locality"><p>Доля Каскелена и Иргели (вкл. Казмис) от общего районного числа по каждому источнику:</p>
<div class="scroll"><table><thead><tr><th>Источник</th><th>Населённый пункт</th><th>Число</th><th>Всего по району</th><th>Доля</th></tr></thead>
<tbody>{loc_html}</tbody></table></div></div>
<div class="panel" id="panel-addr"><p>Областные реестры (без ИИН): привязка <strong>адреса проживания</strong> к учёту ОВД:</p>
<div class="scroll"><table><thead><tr>{"".join(f"<th>{h}</th>" for h in addr_hdr)}</tr></thead>
<tbody>{rows(data["addr_links"], addr_cols)}</tbody></table></div></div>
<div class="panel" id="panel-crime"><p>Подозреваемые, не состоящие на профучёте ОВД. «Место» — место совершения, не проживание:</p>
<div class="scroll"><table><thead><tr><th>ФИО</th><th>ИИН</th><th>ЕРДР</th><th>Статья</th><th>Место совершения</th><th>Район</th></tr></thead>
<tbody>{"".join(f"<tr><td>{esc(r['fio'])}</td><td class='mono'>{esc(r['iin'])}</td><td>{esc(r['erdr'])}</td><td>{esc(r['qual'])}</td><td>{esc(r.get('crime_settlement',''))}</td><td>{esc(r.get('crime_district',''))}</td></tr>" for r in data['suspect_no_police'])}</tbody></table></div></div>
<p style="margin-top:24px"><a href="karasai_spravka.html">← Справка по району</a></p>
</div>
<script>
document.querySelectorAll('.tab').forEach(t=>t.onclick=()=>{{
  document.querySelectorAll('.tab').forEach(x=>x.classList.remove('on'));
  document.querySelectorAll('.panel').forEach(x=>x.classList.remove('on'));
  t.classList.add('on'); document.getElementById('panel-'+t.dataset.tab).classList.add('on');
}});
document.getElementById('search').oninput=e=>{{
  const q=e.target.value.toLowerCase();
  document.querySelectorAll('#main-table tbody tr').forEach(tr=>{{
    tr.style.display=tr.textContent.toLowerCase().includes(q)?'':'none';
  }});
}};
</script></body></html>"""
    path.write_text(html, encoding="utf-8")


def main() -> None:
    data = build_crossmatch()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_html(data, HTML)
    s = data["summary"]
    print(f"Written {OUT}")
    print(f"Written {HTML}")
    print(f"Persons: {s['unique_persons']} | OVD+med: {s['match_ovd_med']} | OVD+med+UD: {s['match_all_three']}")
    print(f"Med gap: {s['gap_med_no_police']} | Crime no police: {s['gap_crime_no_police']} | Addr links: {s['addr_full_linked']}")


if __name__ == "__main__":
    main()
