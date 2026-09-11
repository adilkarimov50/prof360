#!/usr/bin/env python3
"""Нормализация выгрузок диспансера: одна строка заголовков, ИИН без *, автофильтр."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "dispanser_mkb_clean"

SOURCES = [
    Path("/Users/alima_2023/Downloads/A15-19.xlsx"),
    Path("/Users/alima_2023/Downloads/B18-19.xlsx"),
    Path("/Users/alima_2023/Downloads/B20-24.xlsx"),
    Path("/Users/alima_2023/Downloads/D00-09.xlsx"),
    Path("/Users/alima_2023/Downloads/D37-48.xlsx"),
    Path("/Users/alima_2023/Downloads/G30-32.xlsx"),
    Path("/Users/alima_2023/Downloads/G35-37.xlsx"),
    Path("/Users/alima_2023/Downloads/G40.xlsx"),
    Path("/Users/alima_2023/Downloads/I21-22.xlsx"),
    Path("/Users/alima_2023/Downloads/I60-64.xlsx"),
    Path("/Users/alima_2023/Downloads/K74.xlsx"),
    Path("/Users/alima_2023/Downloads/С00-97.xlsx"),
    Path("/Users/alima_2023/Downloads/Орф заб.xlsx"),
]

COLUMNS = [
    "№ п/п",
    "ФИО",
    "ИИН",
    "Дата рождения",
    "Пол",
    "МКБ-10",
    "Диагноз",
    "Организация",
    "Дата постановки на учёт",
    "Система учёта",
]


def full_iin(*cells) -> str:
    for val in cells:
        if val is None or (isinstance(val, float) and pd.isna(val)):
            continue
        s = str(val).strip()
        m = re.search(r"(\d{6})\*+(\d{6})", s)
        if m:
            return m.group(1) + m.group(2)
        digits = re.sub(r"\D", "", s)
        if len(digits) == 12:
            return digits
    return ""


def parse_dob(c3, iin: str) -> str:
    if c3 is not None and not (isinstance(c3, float) and pd.isna(c3)):
        s = str(c3).strip()
        if re.match(r"\d{2}\.\d{2}\.\d{4}", s):
            return s
    if len(iin) == 12:
        yy, mo, d = iin[0:2], iin[2:4], iin[4:6]
        year = f"19{yy}" if int(yy) > 30 else f"20{yy}"
        return f"{d}.{mo}.{year}"
    return ""


def is_org_row(r) -> bool:
    c0, c1, c2 = r[0], r[1], r[2]
    if pd.isna(c0):
        return False
    if pd.isna(c1) and pd.isna(c2) and len(str(c0).strip()) > 20:
        return True
    return False


def parse_file(path: Path) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name=0, header=None)
    rows: list[dict] = []
    org: str | None = None
    ncols = raw.shape[1]

    for i in range(len(raw)):
        r = raw.iloc[i]
        if is_org_row(r):
            org = str(r[0]).strip()
            continue
        c1 = r[1] if ncols > 1 else None
        if pd.isna(c1) or str(c1).strip() in ("", "nan", "Т.А.Ә./Ф.И.О.", "№"):
            continue
        if not str(c1).strip():
            continue
        # строка данных: № в col0 или col0 число
        c0 = r[0]
        if pd.isna(c0):
            continue
        try:
            num = int(float(c0))
        except (ValueError, TypeError):
            continue

        c2 = r[2] if ncols > 2 else None
        c3 = r[3] if ncols > 3 else None
        c4 = r[4] if ncols > 4 else None
        # орф заб: пол в c4, маска полная в c3
        if ncols <= 9 and c4 is not None and str(c4).strip() in ("Муж", "Жен"):
            sex = str(c4).strip()
            icd = r[5] if ncols > 5 else ""
            diag = r[6] if ncols > 6 else ""
            system = r[7] if ncols > 7 else ""
            date_reg = r[8] if ncols > 8 else ""
            iin = full_iin(c2, c3)
            dob = parse_dob(None, iin)
        else:
            sex = r[5] if ncols > 5 else ""
            icd = r[6] if ncols > 6 else ""
            diag = r[7] if ncols > 7 else ""
            system = r[8] if ncols > 8 else ""
            date_reg = r[9] if ncols > 9 else ""
            iin = full_iin(c2, c3, c4)
            dob = parse_dob(c3, iin)

        rows.append(
            {
                "№ п/п": num,
                "ФИО": str(c1).strip(),
                "ИИН": iin,
                "Дата рождения": dob,
                "Пол": str(sex).strip() if pd.notna(sex) else "",
                "МКБ-10": str(icd).strip() if pd.notna(icd) else "",
                "Диагноз": str(diag).strip() if pd.notna(diag) else "",
                "Организация": org or "",
                "Дата постановки на учёт": str(date_reg).strip() if pd.notna(date_reg) else "",
                "Система учёта": str(system).strip() if pd.notna(system) else "",
            }
        )

    return pd.DataFrame(rows, columns=COLUMNS)


def write_excel(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(path, index=False, sheet_name="Данные")
    wb = load_workbook(path)
    ws = wb["Данные"]
    last_row = ws.max_row
    last_col = ws.max_column
    if last_row >= 1 and last_col >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"
        ws.freeze_panes = "A2"
    for col in range(1, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = min(48, 14 if col == 3 else 22)
    wb.save(path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    summary: list[tuple[str, int, int]] = []
    for src in SOURCES:
        if not src.exists():
            print(f"Пропуск (нет файла): {src}")
            continue
        df = parse_file(src)
        without_iin = (df["ИИН"] == "").sum()
        out_name = src.stem + "_clean.xlsx"
        out_path = OUT_DIR / out_name
        write_excel(df, out_path)
        summary.append((src.name, len(df), without_iin))
        print(f"OK {src.name} -> {out_path.name} ({len(df)} строк, без ИИН: {without_iin})")

    pd.DataFrame(summary, columns=["Файл", "Строк", "Без ИИН"]).to_excel(
        OUT_DIR / "_сводка.xlsx", index=False
    )
    print(f"\nПапка: {OUT_DIR}")


if __name__ == "__main__":
    main()
