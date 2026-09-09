"""Генерация прокурорских документов (Word/Excel) с водяными знаками и аудитом."""
import io
from datetime import date, datetime

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.analytics import insights
from app.analytics.scoring import compute_person_factors
from app.core.crypto import decrypt, mask_iin
from app.models.person import Person


def _watermark(username: str, export_id: str) -> str:
    return (f"СЛУЖЕБНО. Сформировано: {username}; {datetime.now().strftime('%d.%m.%Y %H:%M')}; "
            f"ID выгрузки: {export_id}. Запрещено несанкционированное распространение.")


def _docx_base(title: str, username: str, export_id: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = _watermark(username, export_id)
    footer.runs[0].font.size = Pt(7)
    h = doc.add_heading("Прокуратура Алматинской области", level=2)
    doc.add_heading(title, level=1)
    return doc


def _add_kv(doc: Document, key: str, value: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{key}: ").bold = True
    p.add_run(str(value))


def person_report_docx(db: Session, person: Person, username: str, export_id: str,
                       allow_pii: bool) -> bytes:
    doc = _docx_base(f"Справка по лицу: {person.fio}", username, export_id)
    iin = decrypt(person.iin_enc)
    _add_kv(doc, "ИИН", iin if allow_pii else mask_iin(iin))
    _add_kv(doc, "Дата рождения", person.birth_date or "—")
    _add_kv(doc, "Район", person.district or "—")

    score, factors, signals = compute_person_factors(person)
    _add_kv(doc, "Риск-уровень", f"{person.risk_level} ({score} баллов)")

    doc.add_heading("Факторы риска", level=2)
    for f in factors:
        doc.add_paragraph(f"• {f['factor']} (+{f['points']})", style="List Bullet")

    doc.add_heading("Административная практика", level=2)
    for c in sorted(person.admin_cases, key=lambda x: x.case_date or date.min):
        doc.add_paragraph(
            f"{c.case_date or '—'}: {c.qualification or ''} — {(c.measure or 'мера не указана')}. "
            f"Материал № {c.material_no or '—'}.", style="List Bullet")

    if person.preventive_records:
        doc.add_heading("Профилактический учёт", level=2)
        for p in person.preventive_records:
            doc.add_paragraph(
                f"{p.category or p.form}: с {p.date_post or '—'} "
                f"по {p.date_removed or 'наст. время'}.", style="List Bullet")

    if person.suspects:
        doc.add_heading("Сведения о признании подозреваемым (ЕРДР)", level=2)
        for s in person.suspects:
            doc.add_paragraph(
                f"ЕРДР № {s.erdr_no or '—'} ({s.erdr_year or '—'} г.): "
                f"{s.qualification or ''}, {s.gravity or ''}.", style="List Bullet")

    if signals:
        doc.add_heading("Автоматические сигналы", level=2)
        for s in signals:
            doc.add_paragraph(f"[{s['level']}] {s['message']}", style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Документ носит аналитический характер и подлежит проверке прокурором.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def district_report_docx(db: Session, district: str, username: str, export_id: str) -> bytes:
    doc = _docx_base(f"Аналитическая справка по району: {district}", username, export_id)
    counters = insights.overview_counters(db, district)
    _add_kv(doc, "Лиц в районе", counters["persons"])
    _add_kv(doc, "Высокого/критического риска", counters["high_risk"])
    _add_kv(doc, "Подлежат учёту, но не поставлены", counters["should_be_registered"])
    _add_kv(doc, "Повторность на учёте", counters["repeat_on_register"])
    _add_kv(doc, "Эскалация в период учёта", counters["escalation_during_register"])

    doc.add_heading("Топ лиц высокого риска", level=2)
    for p in insights.top_persons(db, 15, district):
        doc.add_paragraph(f"{p['fio']} — {p['risk_level']} ({p['risk_score']})", style="List Bullet")
    doc.add_paragraph()
    doc.add_paragraph("Документ носит аналитический характер и подлежит проверке прокурором.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _act_doc(title: str, addressee: str, username: str, export_id: str) -> Document:
    doc = _docx_base(title, username, export_id)
    p = doc.add_paragraph()
    p.add_run(addressee).bold = True
    return doc


def _legal_footer(doc: Document, used_llm: bool = False) -> None:
    doc.add_paragraph()
    doc.add_paragraph(
        "Правовое основание: Приказ Генерального Прокурора РК от 17.01.2023 № 32 "
        "«О некоторых вопросах организации прокурорского надзора»; Конституционный закон РК «О прокуратуре».")
    note = "Проект акта прокурорского реагирования. Подлежит проверке и подписанию уполномоченным прокурором."
    if used_llm:
        note += " Текст подготовлен с помощью ИИ-помощника и требует юридической выверки."
    doc.add_paragraph(note)


def _render_act_body(doc: Document, text: str) -> None:
    """Рендерит сгенерированный текст акта в DOCX: заголовки разделов и абзацы."""
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        # markdown-заголовки и резолютивные слова -> жирный подзаголовок
        stripped = line.lstrip("#").strip()
        is_heading = (
            line.startswith("#")
            or stripped.upper() in ("ТРЕБУЮ", "ТРЕБУЮ:", "ПРОШУ", "ПРОШУ:")
            or (len(stripped) <= 60 and stripped.endswith(":") and " " not in stripped.rstrip(":"))
        )
        if is_heading:
            p = doc.add_paragraph()
            p.add_run(stripped.rstrip(":")).bold = True
        else:
            doc.add_paragraph(stripped.lstrip("-• ").strip() if stripped[:1] in "-•" else stripped)


def _act_docx(db: Session, kind: str, person: Person, username: str, export_id: str,
              case=None) -> bytes:
    """Общий генератор акта: ИИ-текст (фолбэк — шаблон) + рендер DOCX."""
    from app.ai.act_writer import ACTS, write_act

    meta = ACTS[kind]
    doc = _act_doc(meta["title"], meta["addressee"], username, export_id)
    body, used_llm = write_act(db, kind, person, case)
    _render_act_body(doc, body)
    _legal_footer(doc, used_llm=used_llm)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def representation_docx(db: Session, person: Person, username: str, export_id: str) -> bytes:
    return _act_docx(db, "representation", person, username, export_id)


def protest_docx(db: Session, person: Person, case, username: str, export_id: str) -> bytes:
    return _act_docx(db, "protest", person, username, export_id, case=case)


def appeal_motion_docx(db: Session, person: Person, case, username: str, export_id: str) -> bytes:
    return _act_docx(db, "appeal", person, username, export_id, case=case)


def requirement_docx(db: Session, person: Person, username: str, export_id: str) -> bytes:
    return _act_docx(db, "requirement", person, username, export_id)


def proceedings_violations_xlsx(db: Session, df, username: str, export_id: str) -> bytes:
    """Реестр дел — кандидатов на проверку законности административного производства."""
    from app.analytics import proceedings

    wb = Workbook()
    ws = wb.active
    ws.title = "Адм. производство"
    cols = ["ФИО", "Район", "Материал №", "Дата", "Квалификация",
            "Решение", "Мера", "Признак нарушения (кандидат на проверку)"]
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="922B21")
    for r in proceedings.proceeding_violation_cases(db, df, limit=2000):
        ws.append([r["fio"], r["district"], r["material_no"], r["case_date"],
                   r["qualification"], r["decision"], r["measure"], r["violation"]])

    # сводка по категориям на отдельном листе
    ws2 = wb.create_sheet("Сводка")
    ws2.append(["Категория", "Количество"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="922B21")
    summary = proceedings.admin_proceeding_violations(db, df)
    for cat in summary["categories"]:
        ws2.append([cat["title"], cat["count"]])
    ws2.append(["ИТОГО дел (кандидатов на проверку)", summary["total"]])

    ws.append([])
    ws.append([_watermark(username, export_id)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def escalation_report_xlsx(db: Session, username: str, export_id: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Эскалация в период учёта"
    cols = ["ФИО", "Район", "Категория учёта", "Дата постановки", "Дата снятия",
            "Квалификация УК", "Тяжесть", "Год ЕРДР", "№ ЕРДР"]
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for r in insights.escalation_during_register(db, limit=1000):
        ws.append([r["fio"], r["district"], r["category"], r["date_post"],
                   r["date_removed"] or "состоит", r["qualification"], r["gravity"],
                   r["erdr_year"], r["erdr_no"]])
    ws.append([])
    ws.append([_watermark(username, export_id)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
