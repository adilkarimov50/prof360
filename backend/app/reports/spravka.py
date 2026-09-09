"""Комплексная аналитическая справка с разграничением: полиция (ОВД) и МИО.

Нарратив пишет ИИ (app.ai.spravka_writer) в стиле эталона; статистика и реестр
примеров формируются детерминированно. Поддерживает три вида:
- police   — справка по нарушениям/недостаткам по линии полиции;
- mio      — справка по линии местных исполнительных органов (МИО);
- complex  — общая (комплексная) справка с обоими разделами.
"""
import io

from docx import Document
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.ai import spravka_writer
from app.analytics import insights
from app.reports.generator import _docx_base


def _examples_text(db: Session, df: str | None, limit: int = 12) -> tuple[str, list[dict]]:
    """Текст 'справочно' и список примеров эскалации (ФИО/ЕРДР/статья/дата)."""
    rows = insights.escalation_during_register(db, limit=limit)
    if df:
        rows = [r for r in rows if r.get("district") == df]
    lines = []
    for i, r in enumerate(rows, 1):
        lines.append(
            f"{i}) {r['fio']}, ЕРДР № {r['erdr_no'] or '—'}, {r['qualification'] or '—'}, "
            f"учёт с {r['date_post'] or '—'} (район: {r['district'] or '—'})."
        )
    return ("\n".join(lines) or "Примеры эскалации по выбранному охвату не выявлены."), rows


def _police_facts(db: Session, df: str | None, examples: str) -> str:
    counters = insights.overview_counters(db, df)
    sb = insights.sb_stats(db, df)
    organs = insights.organ_breakdown(db, df)
    top_sb = "; ".join(f"{d['district']} ({d['count']})" for d in sb["by_district"][:6]) or "—"
    return (
        f"Семейно-бытовых протоколов (ст.73 КоАП): {sb['sb_cases']}.\n"
        f"Нарушений защитного предписания (ст.461 КоАП): {sb['zp_violations']}.\n"
        f"Действующих защитных предписаний/особых требований: {sb['protective_orders']}.\n"
        f"Наибольшее число СБ-нарушений: {top_sb}.\n"
        f"Лиц на профилактическом учёте (записей): {counters['preventive']}.\n"
        f"Лиц, подлежащих учёту, но не поставленных: {counters['should_be_registered']}.\n"
        f"Повторность в период учёта: {counters['repeat_on_register']}.\n"
        f"Эскалация в уголовные правонарушения в период учёта: {counters['escalation_during_register']}.\n"
        f"Административных дел по линии полиции (ОВД): {organs['police']}.\n\n"
        f"Справочно (примеры эскалации):\n{examples}"
    )


def _proceedings_facts(db: Session, df: str | None) -> str:
    from app.analytics import proceedings

    summary = proceedings.admin_proceeding_violations(db, df)
    cats = "\n".join(f"- {c['title']}: {c['count']}." for c in summary["categories"])
    cases = proceedings.proceeding_violation_cases(db, df, limit=8)
    examples = "\n".join(
        f"{i}) {c['fio']}, материал № {c['material_no'] or '—'} от {c['case_date'] or '—'}, "
        f"{c['qualification'] or '—'}: {c['decision'] or '—'} / {c['measure']} "
        f"(признак: {c['violation']})."
        for i, c in enumerate(cases, 1)
    ) or "Конкретные примеры по выбранному охвату не выявлены."
    return (
        f"Всего дел — кандидатов на проверку законности административного производства: {summary['total']}.\n"
        f"Распределение по признакам:\n{cats}\n\n"
        f"Справочно (примеры дел, требующих оценки законности постановлений):\n{examples}"
    )


def _mio_facts(db: Session, df: str | None) -> str:
    mio = insights.mio_indicators(db, df)
    organs = insights.organ_breakdown(db, df)
    by_form = "; ".join(f"{f['form']}: {f['count']}" for f in mio["by_form"][:6]) or "—"
    return (
        "Внимание: в исходных данных орган МИО как субъект профилактики отдельно не выделен, "
        "поэтому индикаторы выведены из категорий профилактического учёта (несовершеннолетние, "
        "социально уязвимые формы) и охвата мерами социальной профилактики.\n\n"
        f"Записей профучёта зоны ответственности МИО (соц. профилактика): {mio['mio_preventive_records']}.\n"
        f"Лиц в зоне ответственности МИО: {mio['mio_persons']}.\n"
        f"Из них с признаками эскалации (стали подозреваемыми): {mio['mio_escalated_persons']}.\n"
        f"Разбивка по формам учёта: {by_form}.\n"
        f"Административных дел, отнесённых к линии МИО: {organs['mio']}; не распознано: {organs['unknown']}."
    )


def _add_narrative(doc: Document, scope: str, facts: str) -> bool:
    text, used_llm = spravka_writer.write_section(scope, facts)
    for para in text.split("\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    return used_llm


def _conclusion(doc: Document) -> None:
    doc.add_heading("Выводы и меры прокурорского реагирования", level=2)
    doc.add_paragraph(
        "Изложенное свидетельствует о наличии нарушений законности и формальном подходе к "
        "профилактической работе. По линии прокуратуры (Приказ ГП РК №32 от 17.01.2023) "
        "подлежат рассмотрению следующие акты реагирования:")
    for item in (
        "Представление об устранении нарушений законности — в орган, допустивший бездействие.",
        "Протест — на незаконные/необоснованные постановления по делам об адм. правонарушениях.",
        "Апелляционное ходатайство — по судебным материалам в пределах срока обжалования.",
        "Указание/требование прокурора — о постановке на учёт и усилении профилактического контроля.",
    ):
        doc.add_paragraph(f"• {item}", style="List Bullet")


def _intro(doc: Document, scope_label: str, df: str | None) -> None:
    area = df or "Алматинская область"
    doc.add_paragraph(
        f"Настоящая справка подготовлена службой по защите общественных интересов прокуратуры "
        f"по результатам анализа состояния профилактической работы ({area}). "
        f"Охват: {scope_label}.")


def _build(db: Session, scope: str, df: str | None, username: str, export_id: str) -> bytes:
    titles = {
        "spravka_police": "Аналитическая справка о нарушениях по линии полиции (ОВД)",
        "spravka_mio": "Аналитическая справка о нарушениях по линии МИО",
        "spravka_complex": "Комплексная аналитическая справка о состоянии профилактической работы",
    }
    doc = _docx_base(titles.get(scope, "Аналитическая справка"), username, export_id)

    examples, _rows = _examples_text(db, df)

    if scope in ("spravka_police", "spravka_complex"):
        _intro(doc, "нарушения и недостатки по линии органов внутренних дел", df)
        doc.add_heading("По линии органов внутренних дел (полиции)", level=2)
        _add_narrative(doc, "police", _police_facts(db, df, examples))

        doc.add_heading("Нарушения законности при административном производстве", level=2)
        _add_narrative(doc, "proceedings", _proceedings_facts(db, df))

    if scope in ("spravka_mio", "spravka_complex"):
        if scope == "spravka_mio":
            _intro(doc, "нарушения в действиях местных исполнительных органов (МИО)", df)
        doc.add_heading("По линии местных исполнительных органов (МИО)", level=2)
        _add_narrative(doc, "mio", _mio_facts(db, df))

    _conclusion(doc)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "Документ носит аналитический характер, сформирован с использованием ИИ и подлежит "
        "проверке и подписанию уполномоченным прокурором.")
    note.runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def spravka_police_docx(db: Session, df: str | None, username: str, export_id: str) -> bytes:
    return _build(db, "spravka_police", df, username, export_id)


def spravka_mio_docx(db: Session, df: str | None, username: str, export_id: str) -> bytes:
    return _build(db, "spravka_mio", df, username, export_id)


def spravka_complex_docx(db: Session, df: str | None, username: str, export_id: str) -> bytes:
    return _build(db, "spravka_complex", df, username, export_id)
