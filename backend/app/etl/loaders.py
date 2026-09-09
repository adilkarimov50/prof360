"""Загрузчики исходных Excel-файлов в БД с протоколом загрузки."""
import glob
import os

import openpyxl
import xlrd
from sqlalchemy.orm import Session

from app.analytics.normalize import (
    art_base,
    erdr_year,
    is_special_req,
    norm,
    norm_iin,
    parse_date,
)
from app.core.crypto import iin_hash
from app.etl.column_map import build_admin_map, build_crim_map, row_get
from app.etl.resolver import PersonResolver
from app.models.audit import IngestionLog
from app.models.person import AdminCase, PreventiveRecord, Suspect

# Legacy alias для обратной совместимости
ADM = {
    "organ": 1, "subdivision": 2, "district": 3, "place": 4, "material_no": 5,
    "date": 6, "qual": 9, "fabula": 10, "surname": 11, "name": 12, "patronymic": 13,
    "gender": 14, "res_region": 17, "res_district": 18, "res_locality": 19,
    "iin": 20, "intox": 22, "phone": 23, "decision": 24, "measure": 28, "fine": 29,
}

CRIM = {"erdr": 1, "organ": 3, "iin": 6, "surname": 7, "name": 8, "patronymic": 9,
        "dob": 10, "qual": 11, "gravity": 13, "region": 15}


def _log(db: Session, source_file: str, username: str | None, total: int,
         accepted: int, rejected: int, errors: list) -> IngestionLog:
    entry = IngestionLog(source_file=os.path.basename(source_file), username=username,
                         total=total, accepted=accepted, rejected=rejected,
                         errors={"sample": errors[:20]} if errors else None)
    db.add(entry)
    db.commit()
    return entry


def _read_headers_xlsx(path: str, header_row: int = 2) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers: list[str] = []
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        headers = [norm(c) for c in (row or [])]
        break
    wb.close()
    return headers


def load_admin_excel(db: Session, path: str, source: str = "admin",
                     username: str | None = None, column_map: dict[str, int] | None = None) -> IngestionLog:
    """Загрузка административного массива (источник admin или alcohol)."""
    headers = _read_headers_xlsx(path)
    colmap = column_map or build_admin_map(headers)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    resolver = PersonResolver(db)
    total = accepted = rejected = 0
    errors: list = []

    def g(r, field):
        return row_get(r, colmap, field)

    for r in ws.iter_rows(min_row=3, values_only=True):
        if r is None or all(c is None for c in r):
            continue
        total += 1
        try:
            material_no = norm(g(r, "material_no"))
            if not material_no and not norm(g(r, "surname")):
                rejected += 1
                continue
            iin = norm_iin(g(r, "iin"))
            person = resolver.resolve(
                iin=iin,
                last_name=norm(g(r, "surname")),
                first_name=norm(g(r, "name")),
                patronymic=norm(g(r, "patronymic")),
                gender=norm(g(r, "gender")) or None,
                district=norm(g(r, "district")) or None,
                locality=norm(g(r, "res_locality")) or None,
                address=norm(g(r, "place")) or None,
                phone=norm(g(r, "phone")) or None,
            )
            qual = norm(g(r, "qual"))
            case = AdminCase(
                person_id=person.id,
                material_no=material_no or None,
                case_date=parse_date(g(r, "date")),
                district=norm(g(r, "district")) or None,
                place=norm(g(r, "place")) or None,
                qualification=qual or None,
                article_base=art_base(qual) or None,
                fabula=norm(g(r, "fabula")) or None,
                organ=norm(g(r, "organ")) or None,
                subdivision=norm(g(r, "subdivision")) or None,
                decision=norm(g(r, "decision")) or None,
                measure=norm(g(r, "measure")) or None,
                fine_amount=norm(g(r, "fine")) or None,
                intoxication=norm(g(r, "intox")) or None,
                source=source,
            )
            db.add(case)
            accepted += 1
            if accepted % 1000 == 0:
                db.flush()
        except Exception as e:  # noqa: BLE001
            rejected += 1
            errors.append(f"row {total}: {e}")
    wb.close()
    db.commit()
    return _log(db, path, username, total, accepted, rejected, errors)


def load_criminal_excel(db: Session, path: str, username: str | None = None,
                        column_map: dict[str, int] | None = None) -> IngestionLog:
    """Загрузка реестра подозреваемых (Книга46 / ЕРДР)."""
    headers = _read_headers_xlsx(path, header_row=3)
    colmap = column_map or build_crim_map(headers)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["ЕУСС "] if "ЕУСС " in wb.sheetnames else wb[wb.sheetnames[0]]
    resolver = PersonResolver(db)
    total = accepted = rejected = 0
    errors: list = []

    def g(r, field):
        return row_get(r, colmap, field)

    for r in ws.iter_rows(min_row=4, values_only=True):
        if r is None or all(c is None for c in r):
            continue
        total += 1
        try:
            erdr = norm(g(r, "erdr"))
            iin = norm_iin(g(r, "iin"))
            ih = iin_hash(iin)
            person = None
            if iin:
                person = resolver.resolve(
                    iin=iin,
                    last_name=norm(g(r, "surname")),
                    first_name=norm(g(r, "name")),
                    patronymic=norm(g(r, "patronymic")),
                    birth_date=parse_date(g(r, "dob")),
                    district=norm(g(r, "region")) or None,
                )
            qual = norm(g(r, "qual"))
            db.add(Suspect(
                person_id=person.id if person else None,
                erdr_no=erdr or None,
                erdr_year=erdr_year(erdr),
                qualification=qual or None,
                gravity=norm(g(r, "gravity")) or None,
                organ=norm(g(r, "organ")) or None,
                region=norm(g(r, "region")) or None,
                iin_hash=ih,
            ))
            accepted += 1
        except Exception as e:  # noqa: BLE001
            rejected += 1
            errors.append(f"row {total}: {e}")
    wb.close()
    db.commit()
    return _log(db, path, username, total, accepted, rejected, errors)


def load_preventive_xls(db: Session, path: str, username: str | None = None) -> IngestionLog:
    """Загрузка профучёта (.xls формы 205/206/301/УДО). Маппинг по названиям колонок."""
    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    headers = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
    resolver = PersonResolver(db)
    form = os.path.basename(path).split("_")[0]

    def col(*names: str) -> int | None:
        for n in names:
            for i, h in enumerate(headers):
                if n.lower() in h.lower():
                    return i
        return None

    c_last = col("Фамилия")
    c_first = col("Имя")
    c_patr = col("Отчество")
    c_dob = col("Дата рождения")
    c_iin = col("ИИН")
    c_district = col("Район")
    c_locality = col("Населенный пункт")
    c_status = col("Статус")
    c_date = col("Дата ввода")

    total = accepted = rejected = 0
    errors: list = []
    for rr in range(1, sheet.nrows):
        total += 1
        try:
            def cell(idx):
                return sheet.cell_value(rr, idx) if idx is not None else ""

            iin = norm_iin(cell(c_iin))
            status = norm(cell(c_status))
            person = resolver.resolve(
                iin=iin,
                last_name=norm(cell(c_last)),
                first_name=norm(cell(c_first)),
                patronymic=norm(cell(c_patr)),
                birth_date=parse_date(cell(c_dob)),
                district=norm(cell(c_district)) or None,
                locality=norm(cell(c_locality)) or None,
            )
            category = f"Форма {form}. {status}".strip()
            db.add(PreventiveRecord(
                person_id=person.id,
                form=form,
                category=category,
                status=status or None,
                date_post=parse_date(cell(c_date)),
                date_removed=None,
                district=norm(cell(c_district)) or None,
                locality=norm(cell(c_locality)) or None,
                has_special_req=is_special_req(category),
            ))
            accepted += 1
        except Exception as e:  # noqa: BLE001
            rejected += 1
            errors.append(f"row {total}: {e}")
    db.commit()
    return _log(db, path, username, total, accepted, rejected, errors)


def discover_and_load_all(db: Session, data_dir: str, username: str | None = None) -> list[dict]:
    """Находит исходные файлы в каталоге и загружает их. Возвращает сводку."""
    summary: list[dict] = []

    for path in sorted(glob.glob(os.path.join(data_dir, "*.xlsx"))):
        base = os.path.basename(path)
        if base.startswith("~$"):
            continue
        if "алк" in base.lower():
            log = load_admin_excel(db, path, source="alcohol", username=username)
        elif base.startswith("Книга46") or "ердр" in base.lower():
            log = load_criminal_excel(db, path, username=username)
        elif base.startswith("ТЗ") or base.startswith("Реестр") or base.startswith("Справка"):
            continue
        else:
            log = load_admin_excel(db, path, source="admin", username=username)
        summary.append({"file": base, "accepted": log.accepted, "rejected": log.rejected})

    prof_dir = os.path.join(data_dir, "проф дела")
    for path in sorted(glob.glob(os.path.join(prof_dir, "*.xls"))):
        log = load_preventive_xls(db, path, username=username)
        summary.append({"file": os.path.basename(path), "accepted": log.accepted, "rejected": log.rejected})

    return summary


def preview_excel_columns(path: str) -> dict:
    """Возвращает заголовки и предложенный маппинг для UI/ИИ."""
    base = os.path.basename(path).lower()
    if path.endswith(".xls"):
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        headers = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
        return {"file": os.path.basename(path), "type": "preventive", "headers": headers}
    headers = _read_headers_xlsx(path, header_row=3 if "книга46" in base or "ердр" in base else 2)
    if "книга46" in base or "ердр" in base:
        return {
            "file": os.path.basename(path), "type": "criminal",
            "headers": headers, "suggested_map": build_crim_map(headers),
        }
    return {
        "file": os.path.basename(path),
        "type": "alcohol" if "алк" in base else "admin",
        "headers": headers,
        "suggested_map": build_admin_map(headers),
    }
