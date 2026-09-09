"""Извлечение текста из документов комиссии."""
from __future__ import annotations

import io

from docx import Document

from app.ai import llm

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def extract_docx_text(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_text(file_bytes: bytes, mime_type: str, filename: str = "") -> str:
    """Извлечь текст: .docx локально, PDF/скан через Gemini multimodal."""
    mt = (mime_type or "").lower()
    ext = (filename or "").lower().rsplit(".", 1)[-1] if filename else ""

    if mt == DOCX_MIME or ext == "docx":
        text = extract_docx_text(file_bytes)
        if text:
            return text
        raise ValueError("Не удалось извлечь текст из DOCX")

    if mt in llm.MULTIMODAL_MIMES or ext in ("pdf", "jpg", "jpeg", "png"):
        resolved = mt
        if not resolved or resolved == "application/octet-stream":
            resolved = {
                "pdf": "application/pdf",
                "jpg": "image/jpeg",
                "jpeg": "image/jpeg",
                "png": "image/png",
            }.get(ext, mt)
        text = llm.extract_text_multimodal(file_bytes, resolved)
        if text:
            return text.strip()
        raise ValueError("Gemini multimodal недоступен или не распознал документ")

    raise ValueError(
        f"Неподдерживаемый формат ({mime_type}). "
        "Поддерживаются PDF, DOCX, JPG, PNG. Старый .doc конвертируйте в PDF или DOCX."
    )
